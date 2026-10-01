# -*- coding: utf-8 -*-
"""아주대 의대 2026 해부학 수업 시간표 xlsx → 앱 페이로드 변환 (의예과 2학년).

포맷: 행 = 날짜, 열 = 1~7교시 (09:00~17:00, 60분, 점심 12~13시 — config.js SCHOOLS.ajou.periods).
셀: '팔 강의(김한기)', '등 실습', '등-머리\\n구연시험', '자율학습1' 등. 줄바꿈은 셀 폭 때문에
들어간 것이라 공백으로 합치되, '피드\\n백'처럼 한 글자만 넘어간 경우는 붙인다.
과목명 = 셀 내용(부위별: '등 실습', '팔 강의' …), 괄호 안 = 교수.
같은 문구가 한 셀에 3번 반복된 칸(12/10 이론시험·12/11 실습시험)은 원본에서 3교시짜리
병합이 풀린 흔적 — 그 칸을 가운데로 앞뒤 1교시씩 펼친다.

2026-10-01 수정본(2026년_해부학_수업_시간표.xlsx): 줄바꿈 정리·'오전만 운영' 칸·12/8 구연시험 7교시·
12/11 실습시험 2~4교시로 확정. 3회 반복 셀은 없음(구버전 호환용으로 처리 유지).

사용: python3 tools/build_ajou_2026_2.py <xlsx경로> <출력디렉토리>
"""
import sys, os, re, json, time
from datetime import date, datetime
import openpyxl

PERIOD_START = {1:'9:00',2:'10:00',3:'11:00',4:'13:00',5:'14:00',6:'15:00',7:'16:00'}
PERIOD_END   = {1:'10:00',2:'11:00',3:'12:00',4:'14:00',5:'15:00',6:'16:00',7:'17:00'}
WKN = ['월','화','수','목','금','토','일']
GRADE = '의예과 2학년'

def clean(raw):
    lines = [l.strip() for l in str(raw).split('\n') if l.strip()]
    out = ''
    for l in lines:
        # 한 글자(+괄호)만 다음 줄로 넘어간 경우: '피드' + '백(박진서)'
        if out and re.match(r'^[가-힣](\(|$)', l):
            out += l
        else:
            out = (out + ' ' + l) if out else l
    out = re.sub(r'\s*,\s*', ', ', out)
    return re.sub(r'\s+', ' ', out).strip()

def dedupe_repeat(s):
    """'A A A' 처럼 같은 문구가 반복되면 (문구, 반복수)"""
    words = s.split(' ')
    for n in range(1, len(words) // 2 + 1):
        if len(words) % n:
            continue
        unit = words[:n]
        if unit * (len(words) // n) == words:
            return ' '.join(unit), len(words) // n
    return s, 1

def parse_cell(raw):
    s = clean(raw)
    s, rep = dedupe_repeat(s)
    prof = ''
    m = re.match(r'^(.*?)\s*\(([^()]+)\)\s*$', s)
    if m:
        s, prof = m.group(1).strip(), m.group(2).strip()
    s = re.sub(r'^(자율학습)\d+$', r'\1', s)
    s = re.sub(r'(이론|실습|구연)\s+시험', r'\1시험', s)
    return s, prof, rep

def main(xlsx_path, out_dir):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]
    # 헤더 행: '날짜' 열 + 'N교시' 열
    hdr_row, pcols = None, {}
    for r in range(1, 10):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(str(v).strip() == '날짜' for v in vals if v):
            hdr_row = r
            for c, v in enumerate(vals, 1):
                pm = re.match(r'^\s*(\d+)\s*교시', str(v or ''))
                if pm:
                    pcols[c] = int(pm.group(1))
            break
    assert hdr_row and pcols, '헤더 행을 찾지 못함'

    slots = {}  # (date, period) -> item
    start_date = None
    for r in range(hdr_row + 1, ws.max_row + 1):
        dv = ws.cell(r, 1).value
        if dv is None:
            continue
        ds = dv.strftime('%Y-%m-%d') if isinstance(dv, (datetime, date)) else str(dv).strip()[:10]
        if not re.match(r'^\d{4}-\d{2}-\d{2}$', ds):
            continue
        d = date.fromisoformat(ds)
        start_date = start_date or d
        for c, p in pcols.items():
            v = ws.cell(r, c).value
            if v is None or not str(v).strip():
                continue
            subj, prof, rep = parse_cell(v)
            if re.fullmatch(r'오전만\s*운영', subj):  # 11/2~11/6 오후 = 수업 없음 표시
                continue
            ps = [p]
            if rep == 3:
                ps = [q for q in (p - 1, p, p + 1) if q in PERIOD_START]
            for q in ps:
                slots[(ds, q)] = (d, subj, prof)

    items, wdd, ed = [], {}, set()
    monday0 = start_date.toordinal() - start_date.weekday()
    for (ds, p) in sorted(slots):
        d, subj, prof = slots[(ds, p)]
        week = str((d.toordinal() - monday0) // 7 + 1)
        day = WKN[d.weekday()]
        is_exam = '시험' in subj
        items.append({'week': week, 'date': ds, 'day': day, 'period': p,
                      'start': PERIOD_START[p], 'end': PERIOD_END[p],
                      'subject': subj, 'professor': prof, 'is_exam': is_exam})
        wdd.setdefault(week, {})[day] = ds
        if is_exam:
            ed.add(ds)
    # 수업 없는 평일도 날짜 칸이 비지 않게 채움
    for wk, dd in wdd.items():
        any_ds = date.fromisoformat(next(iter(dd.values())))
        mon = any_ds.toordinal() - any_ds.weekday()
        for i in range(5):
            dd.setdefault(WKN[i], date.fromordinal(mon + i).isoformat())

    wks = sorted(wdd, key=int)
    ts = int(time.time() * 1000)
    payload = {'items': items, 'wdd': {k: wdd[k] for k in wks}, 'ed': sorted(ed), 'wks': wks,
               'grade': GRADE, 'ts': ts,
               'changelog': {'ts': ts, 'msg': '2026 해부학 수업 시간표 적용'}}
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, 'ajou_' + GRADE.replace(' ', '_') + '.json')
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    print(out, len(items), '항목', '주차', wks[0], '~', wks[-1], '시험일', sorted(ed))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
